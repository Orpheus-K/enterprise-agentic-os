import io
import json
import zipfile

import pytest
from docx import Document
from pypdf import PdfWriter

from agentos_compute.documents import DOCX, MAX_INPUT, PDF, DocumentError, extract


def active() -> None:
    pass


def docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("中文合同：合成测试数据")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "金额 100"
    document.add_paragraph("结束")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def test_chinese_docx() -> None:
    result = json.loads(extract(docx_bytes(), DOCX, active))
    assert result == {
        "format": "agentos.document-text.v1",
        "blocks": ["中文合同：合成测试数据", "金额 100", "结束"],
    }


def test_blank_pdf_and_encryption() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    output = io.BytesIO()
    writer.write(output)
    assert json.loads(extract(output.getvalue(), PDF, active))["blocks"] == [""]
    writer.encrypt("synthetic")
    output = io.BytesIO()
    writer.write(output)
    with pytest.raises(DocumentError, match="encrypted_document"):
        extract(output.getvalue(), PDF, active)


@pytest.mark.parametrize(
    ("data", "media", "reason"),
    [
        (b"x", PDF, "malformed_document"),
        (b"x", DOCX, "malformed_document"),
        (b"x", "text/html", "unsupported_media_type"),
        (b"x" * (MAX_INPUT + 1), PDF, "input_size"),
    ],
    ids=["bad-pdf", "bad-docx", "unsupported", "oversize"],
)
def test_rejections(data: bytes, media: str, reason: str) -> None:
    with pytest.raises(DocumentError, match=reason):
        extract(data, media, active)


@pytest.mark.parametrize(
    ("name", "content", "reason"),
    [
        ("../outside", b"x", "unsafe_archive"),
        ("word/vbaProject.bin", b"x", "unsafe_archive"),
        ("bad.xml", b'<!DOCTYPE a [<!ENTITY x "x">]><a/>', "unsafe_xml"),
        (
            "bad.xml",
            '<!DOCTYPE a [<!ENTITY x "x">]><a/>'.encode("utf-16"),
            "unsupported_xml_encoding",
        ),
        (
            "bad.rels",
            b'<Relationships><r TargetMode="External"/></Relationships>',
            "external_relationship",
        ),
    ],
)
def test_unsafe_archive(name: str, content: bytes, reason: str) -> None:
    output = io.BytesIO(docx_bytes())
    with zipfile.ZipFile(output, "a") as archive:
        archive.writestr(name, content)
    with pytest.raises(DocumentError, match=reason):
        extract(output.getvalue(), DOCX, active)


def test_cancel_during_archive_validation() -> None:
    calls = 0

    def cancelled() -> None:
        nonlocal calls
        calls += 1
        if calls == 3:
            raise TimeoutError("cancelled")

    with pytest.raises(TimeoutError, match="cancelled"):
        extract(docx_bytes(), DOCX, cancelled)
